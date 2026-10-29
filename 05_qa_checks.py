"""Step 5 - Independent QA: recompute every KPI in Python and reconcile it with the recalculated Excel workbook.
Usage: python scripts/05_qa_checks.py [path_to_recalculated_xlsx]"""
import sys, json, pathlib, math
import openpyxl
ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "dashboard"))
from kpi_engine import load, compute_kpis, by_group
xl = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "excel/operations_analysis.xlsx"
M = json.load(open(ROOT / "scripts/_workbook_map.json"))
wb = openpyxl.load_workbook(xl, data_only=True)
d = load(ROOT / "data/processed/cleaned_orders.csv"); k = compute_kpis(d)
fails, n = [], 0
def cmp(label, a, b, tol=1e-6):
    global n; n += 1
    if a is None or b is None or isinstance(a, (str, list)) or isinstance(b, (str, list)):
        ok = a == b
    else:
        ok = math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=tol)
    if not ok: fails.append((label, a, b))
ws = wb["KPI ANALYSIS"]
for key, ref in M["KREF"].items():
    cell = ref.split("!")[1].replace("$", "")
    cmp(f"KPI {key}", ws[cell].value, k[key])
# pivot by category / region / customer type
cols = {"Orders": "B", "Completed": "C", "Pending": "D", "Cancelled": "E", "On-time delivery %": "H", "Late deliveries": "I",
        "Avg delay when late (d)": "J", "On-time install %": "K", "Avg production (d)": "L", "Avg cycle (d)": "M", "Overdue": "N"}
pv = wb["PIVOT ANALYSIS"]
for key, dim in [("cat", "Product_Category"), ("region", "Region"), ("ctype", "Customer_Type")]:
    g = by_group(d, dim).set_index(dim)
    r = M["PIV"][key]
    while pv[f"A{r}"].value not in (None, "Overall"):
        lab = pv[f"A{r}"].value
        for h, c in cols.items():
            v = pv[f"{c}{r}"].value; exp = g.loc[lab, h]
            if v == "" or v is None: v = None
            exp = None if exp is None or (isinstance(exp, float) and math.isnan(exp)) else exp
            cmp(f"PIVOT {dim}={lab} {h}", v, exp)
        r += 1
# dashboard cards with filter All
db = wb["OPERATIONS DASHBOARD"]
cmp("DASH total", db["B8"].value, k["total_orders"]); cmp("DASH completion", db["D8"].value, k["order_completion_rate"])
cmp("DASH backlog", db["F8"].value, k["backlog"]); cmp("DASH on-time del", db["H8"].value, k["on_time_delivery_rate"])
cmp("DASH install completion", db["J8"].value, k["installation_completion_rate"]); cmp("DASH on-time install", db["L8"].value, k["on_time_install_rate"])
cmp("DASH overdue", db["N8"].value, k["overdue_orders"]); cmp("DASH value Cr", db["P8"].value, k["total_order_value"] / 1e7)
filt = db["C4"].value, db["H4"].value, db["M4"].value
# overdue list
ov = d[d.Overdue == 1].sort_values(["Days_Overdue"], ascending=False)
if filt == ("All", "All", "All"):
    xl_ids = [db[f"B{r}"].value for r in range(59, 69) if db[f"B{r}"].value]
    py = d[d.Overdue == 1].assign(key=lambda x: x.Days_Overdue.astype(float) + (x.index + 2) / 100000).sort_values("key", ascending=False).Order_ID.tolist()[:10]
    cmp("DASH overdue list", xl_ids, py)
# data quality live checks
dq = wb["DATA QUALITY LOG"]
for r in range(6, 17):
    if dq[f"E{r}"].value == "CHECK": fails.append((f"DQ check {dq[f'B{r}'].value}", dq[f"C{r}"].value, "expected OK"))
# formula errors anywhere
errs = 0
for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("#") and c.value.rstrip("!?0/").upper() in ("#REF", "#VALUE", "#DIV", "#NAME", "#N/A", "#NUM"):
                errs += 1
# independent hand-style recount of headline KPIs straight from the raw CSV (no shared code)
import csv, datetime as dt
raw = list(csv.DictReader(open(ROOT / "data/raw/raw_orders_export.csv", encoding="utf-8")))
seen, uniq = set(), []
for r_ in raw:
    t = tuple(v.strip() for v in r_.values())
    if t not in seen: seen.add(t); uniq.append(r_)
st = [r_["Status"].strip().lower() for r_ in uniq]
cmp("RAW recount total", len(uniq), k["total_orders"])
cmp("RAW recount completed", sum(s_ in ("completed", "complete") for s_ in st), k["completed"])
cmp("RAW recount cancelled", sum(s_ in ("cancelled", "canceled") for s_ in st), k["cancelled"])
print(f"checks run: {n}, failures: {len(fails)}, error cells: {errs}")
for f_ in fails[:40]: print("FAIL", f_)
print("dashboard filter state:", filt)
sys.exit(1 if fails or errs else 0)
