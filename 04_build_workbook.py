"""
Step 4 - Build the Excel workbook  excel/operations_analysis.xlsx
Every KPI, summary table, chart series and dashboard card is an Excel FORMULA that reads the
CLEANED DATA sheet, so the workbook recalculates if the data or the report date changes.
Narrative text (findings) quotes numbers from analysis/results.json, which the QA script
reconciles against the recalculated workbook.
"""
import json, sys, pathlib, csv
from datetime import datetime, date
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
def DLL():
    return DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False)
from openpyxl.chart.layout import Layout, ManualLayout

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts")); sys.path.insert(0, str(ROOT / "dashboard"))
from data_dictionary import RAW as DD_RAW, CLEAN as DD_CLEAN, EXCEL_ONLY as DD_X
from kpi_engine import INPUT_COLS
R = json.load(open(ROOT / "analysis/results.json"))
K = R["kpis"]

# ---------- style ----------
TEAL, CHAR, SILVER, GREEN, AMBER, RED, MID = "16B8B5", "2D3136", "F4F6F8", "3AAE4B", "E8A33D", "D9534F", "8A939C"
F = "Arial"
def font(sz=10, b=False, c=CHAR, i=False): return Font(name=F, size=sz, bold=b, color=c, italic=i)
def fill(c): return PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="D5DADF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
PCT, INT, DEC, INR, DT = "0.0%", "#,##0", "0.0", "₹#,##0", "dd-mmm-yyyy"

wb = Workbook()
ws0 = wb.active; ws0.title = "START HERE"
names = ["RAW DATA", "CLEANED DATA", "KPI ANALYSIS", "PIVOT ANALYSIS", "OPERATIONS DASHBOARD",
         "BUSINESS RECOMMENDATIONS", "DATA QUALITY LOG", "DATA DICTIONARY", "LISTS"]
S = {n: wb.create_sheet(n) for n in names}
S["START HERE"] = ws0
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False

def header_row(ws, row, labels, col=1, color=CHAR):
    for j, h in enumerate(labels):
        c = ws.cell(row, col + j, h); c.font = font(10, True, "FFFFFF"); c.fill = fill(color)
        c.alignment = CENTER; c.border = BOX

def title(ws, text, sub=None, width_cols=10):
    ws["A1"] = text; ws["A1"].font = font(16, True, CHAR)
    ws.row_dimensions[1].height = 26
    if sub:
        ws["A2"] = sub; ws["A2"].font = font(9, False, RED, True)

SYN = "SIMULATED DATA - synthetic records created for a portfolio case study. Not real company performance."

# ======================= LISTS (hidden) =======================
CATS = ["Heat Pump Water Heater", "Commercial Heat Pump", "Swimming Pool Heat Pump", "Heat Pump Dryer", "Custom Requirement"]
CTYPES = ["Residential", "Commercial", "Institutional"]
REGIONS = ["Bengaluru", "Rest of Karnataka", "Other States"]
L = S["LISTS"]
for j, (h, vals) in enumerate([("Category", ["All"] + CATS), ("Customer", ["All"] + CTYPES), ("Region", ["All"] + REGIONS),
                               ("Status", ["Completed", "Pending", "Cancelled"])], start=1):
    L.cell(1, j, h)
    for i, v in enumerate(vals, start=2): L.cell(i, j, v)
L.sheet_state = "hidden"

# ======================= RAW DATA =======================
ws = S["RAW DATA"]
title(ws, "RAW DATA - original export exactly as received (do not edit)", SYN)
ws["A3"] = "163 rows including 3 duplicate rows and deliberate data-entry problems. Values are stored as text, as an export would deliver them. The cleaned version is on CLEANED DATA."
ws["A3"].font = font(9, i=True)
with open(ROOT / "data/raw/raw_orders_export.csv", encoding="utf-8") as f:
    rows = list(csv.reader(f))
header_row(ws, 5, rows[0])
for i, r in enumerate(rows[1:], start=6):
    for j, v in enumerate(r, start=1):
        c = ws.cell(i, j, v if v != "" else None); c.font = font(9); c.number_format = "@"
for j, w in enumerate([12, 12, 12, 13, 16, 24, 7, 15, 13, 13, 13, 14, 13, 12, 22, 24, 24, 34], start=1):
    ws.column_dimensions[CL(j)].width = w
ws.freeze_panes = "B6"; ws.auto_filter.ref = f"A5:R{5 + len(rows) - 1}"

# ======================= CLEANED DATA =======================
ws = S["CLEANED DATA"]
d = pd.read_csv(ROOT / "data/processed/cleaned_orders.csv", parse_dates=["Order_Date", "Expected_Delivery_Date", "Dispatch_Date",
    "Actual_Delivery_Date", "Expected_Installation_Date", "Actual_Installation_Date"], usecols=INPUT_COLS)
DERIVED = ["Order_Month", "Delivery_Due", "Delivered", "Delivery_Delay_Days", "Delivery_Timeliness", "Install_Due", "Installed",
           "Install_Delay_Days", "Install_Timeliness", "Production_Days", "Transit_Days", "Install_Lead_Days",
           "Order_To_Delivery_Days", "Order_Cycle_Days", "Current_Stage", "Overdue", "Days_Overdue", "Overdue_Stage",
           "In_Dashboard_Filter", "Overdue_Rank_Key"]
COLS = INPUT_COLS + DERIVED
COL = {c: CL(i + 1) for i, c in enumerate(COLS)}
LAST = 1000  # formulas cover rows 2-1000 so new orders can be appended
def RG(c): return f"'CLEANED DATA'!${COL[c]}$2:${COL[c]}${LAST}"
T = {
 "Order_Month": '=DATE(YEAR({Order_Date}),MONTH({Order_Date}),1)',
 "Delivery_Due": '=IF(AND({Order_Status}<>"Cancelled",{Expected_Delivery_Date}<=AsOf),1,0)',
 "Delivered": '=IF(AND({Order_Status}<>"Cancelled",OR(ISNUMBER({Actual_Delivery_Date}),{Order_Status}="Completed")),1,0)',
 "Delivery_Delay_Days": '=IF(AND({Delivered}=1,{Valid_Delivery_Timing}=1,ISNUMBER({Actual_Delivery_Date})),{Actual_Delivery_Date}-{Expected_Delivery_Date},"")',
 "Delivery_Timeliness": '=IF({Delivery_Delay_Days}="","",IF({Delivery_Delay_Days}<0,"Early",IF({Delivery_Delay_Days}=0,"On Time","Late")))',
 "Install_Due": '=IF(AND({Order_Status}<>"Cancelled",{Expected_Installation_Date}<=AsOf),1,0)',
 "Installed": '=IF(AND({Order_Status}<>"Cancelled",OR(ISNUMBER({Actual_Installation_Date}),{Order_Status}="Completed")),1,0)',
 "Install_Delay_Days": '=IF(AND({Installed}=1,{Valid_Install_Timing}=1,ISNUMBER({Actual_Installation_Date})),{Actual_Installation_Date}-{Expected_Installation_Date},"")',
 "Install_Timeliness": '=IF({Install_Delay_Days}="","",IF({Install_Delay_Days}<0,"Early",IF({Install_Delay_Days}=0,"On Time","Late")))',
 "Production_Days": '=IF(AND({Valid_Dispatch_Timing}=1,ISNUMBER({Dispatch_Date}),{Order_Status}<>"Cancelled"),{Dispatch_Date}-{Order_Date},"")',
 "Transit_Days": '=IF(AND({Production_Days}<>"",{Delivery_Delay_Days}<>""),{Actual_Delivery_Date}-{Dispatch_Date},"")',
 "Install_Lead_Days": '=IF(AND({Delivery_Delay_Days}<>"",{Install_Delay_Days}<>""),{Actual_Installation_Date}-{Actual_Delivery_Date},"")',
 "Order_To_Delivery_Days": '=IF({Delivery_Delay_Days}<>"",{Actual_Delivery_Date}-{Order_Date},"")',
 "Order_Cycle_Days": '=IF(AND({Order_Status}="Completed",{Install_Delay_Days}<>""),{Actual_Installation_Date}-{Order_Date},"")',
 "Current_Stage": '=IF({Order_Status}="Cancelled","Cancelled",IF({Order_Status}="Completed","Completed",IF({Delivered}=1,"Awaiting Installation",IF(ISNUMBER({Dispatch_Date}),"In Transit","In Production"))))',
 "Overdue": '=IF(AND({Order_Status}="Pending",OR(AND({Delivered}=0,{Expected_Delivery_Date}<AsOf),AND({Delivered}=1,{Expected_Installation_Date}<AsOf))),1,0)',
 "Days_Overdue": '=IF({Overdue}=1,IF({Delivered}=0,AsOf-{Expected_Delivery_Date},AsOf-{Expected_Installation_Date}),"")',
 "Overdue_Stage": '=IF({Overdue}=1,IF({Delivered}=0,"Delivery","Installation"),"")',
 "In_Dashboard_Filter": '=IF(AND(OR(fCat="All",{Product_Category}=fCat),OR(fCust="All",{Customer_Type}=fCust),OR(fReg="All",{Region}=fReg)),1,0)',
 "Overdue_Rank_Key": '=IF(AND({Overdue}=1,{In_Dashboard_Filter}=1),{Days_Overdue}+ROW()/100000,0)',
}
header_row(ws, 1, COLS)
for j, c in enumerate(COLS, start=1):
    ws.cell(1, j).fill = fill(TEAL if c in DERIVED else CHAR)
DATECOLS = {"Order_Date", "Expected_Delivery_Date", "Dispatch_Date", "Actual_Delivery_Date", "Expected_Installation_Date", "Actual_Installation_Date", "Order_Month"}
for i, rec in enumerate(d.itertuples(index=False), start=2):
    for j, c in enumerate(INPUT_COLS, start=1):
        v = getattr(rec, c)
        if isinstance(v, float) and pd.isna(v): v = None
        elif v is pd.NaT or (hasattr(v, "year") and pd.isna(v)): v = None
        elif isinstance(v, pd.Timestamp): v = v.to_pydatetime()
        elif isinstance(v, str) and v.strip() == "": v = None
        cell = ws.cell(i, j, v); cell.font = font(9)
        if c in DATECOLS: cell.number_format = DT
        if c == "Order_Value_INR": cell.number_format = INR
    for c in DERIVED:
        f = T[c]
        for k2, letter in COL.items():
            f = f.replace("{" + k2 + "}", f"{letter}{i}")
        cell = ws.cell(i, COLS.index(c) + 1, f); cell.font = font(9, c=CHAR)
        if c in DATECOLS: cell.number_format = "mmm-yyyy"
NROWS = len(d) + 1
for j, c in enumerate(COLS, start=1):
    ws.column_dimensions[CL(j)].width = max(11, min(26, len(c) + 2))
ws.column_dimensions[COL["Operational_Notes"]].width = 30
ws.column_dimensions[COL["DQ_Flag"]].width = 30
ws.freeze_panes = "B2"; ws.auto_filter.ref = f"A1:{CL(len(COLS))}{NROWS}"
ws.conditional_formatting.add(f"{COL['DQ_Flag']}2:{COL['DQ_Flag']}{NROWS}",
    CellIsRule(operator="notEqual", formula=['"OK"'], fill=fill("FDECEA"), font=Font(name=F, size=9, color=RED)))
ws.conditional_formatting.add(f"{COL['Overdue']}2:{COL['Overdue']}{NROWS}",
    CellIsRule(operator="equal", formula=["1"], fill=fill("FDECEA"), font=Font(name=F, size=9, bold=True, color=RED)))
dv = DataValidation(type="list", formula1="=LISTS!$D$2:$D$4", allow_blank=False)
ws.add_data_validation(dv); dv.add(f"{COL['Order_Status']}2:{COL['Order_Status']}{LAST}")
ws.cell(1, len(COLS) + 2, "Dark headers = validated input values. Teal headers = formulas (do not type over).").font = font(9, i=True)

# ======================= KPI ANALYSIS =======================
ws = S["KPI ANALYSIS"]
title(ws, "KPI ANALYSIS - all results are live formulas on CLEANED DATA", SYN)
ws["B4"] = "Report as-of date  ►"; ws["B4"].font = font(10, True); ws["B4"].alignment = Alignment(horizontal="right")
ws["C4"] = datetime(2026, 6, 30); ws["C4"].number_format = DT; ws["C4"].font = font(11, True, "0000FF")
ws["C4"].fill = fill("FFF59D"); ws["C4"].border = BOX
ws["D4"] = "Yellow = editable input. End of the Jan-Jun 2026 review period. 'Due' and 'overdue' are judged against this date, never against today's date."
ws["D4"].font = font(9, i=True)
wb.defined_names["AsOf"] = DefinedName("AsOf", attr_text="'KPI ANALYSIS'!$C$4")
HDR = ["#", "KPI", "Result", "Unit", "Numerator", "Denominator", "Formula (plain English)", "Business meaning",
       "Applicable records / exclusions", "Interpretation (data as of 30-Jun-2026)"]
header_row(ws, 6, HDR)
st, A, = RG("Order_Status"), RG("Order_ID")
def pct(n, dn): return f"=IFERROR({n}/{dn},\"n/a\")"
KREF = {}
NUM = [0]
row = 7
def section(name):
    global row
    ws.cell(row, 1, name).font = font(11, True, "FFFFFF")
    for j in range(1, 11): ws.cell(row, j).fill = fill(TEAL)
    row += 1
def kpi(key, name, res, unit, num, den, plain, meaning, applies, interp):
    global row
    vals = [None, name, res, unit, num, den, plain, meaning, applies, interp]
    NUM[0] += 1
    ws.cell(row, 1, NUM[0] if unit != "check" else "✓")
    if unit == "check": NUM[0] -= 1
    for j, v in enumerate(vals[1:], start=2):
        c = ws.cell(row, j, v); c.alignment = WRAP; c.font = font(9); c.border = BOX
    ws.cell(row, 1).border = BOX; ws.cell(row, 1).font = font(9); ws.cell(row, 1).alignment = WRAP
    ws.cell(row, 2).font = font(9, True); ws.cell(row, 3).font = font(10, True, TEAL if unit == "%" else CHAR)
    fmt = {"%": PCT, "count": INT, "days": DEC, "INR": INR, "check": "@"}[unit]
    ws.cell(row, 3).number_format = fmt
    for j in (5, 6): ws.cell(row, j).number_format = INR if unit == "INR" else (INT if unit != "days" else "0.0")
    KREF[key] = f"'KPI ANALYSIS'!$C${row}"
    row += 1
def ref(key, col="C"): return KREF[key].replace("$C$", f"${col}$")
def p(x): return f"{x*100:.1f}%"

section("A. ORDER MANAGEMENT")
kpi("total_orders", "Total Orders", f"=COUNTA({A})", "count", None, None, "Count of unique Order IDs", "Size of the workload in the period", "All 160 cleaned orders (3 duplicate rows already removed)", f"{K['total_orders']} orders booked Jan-Jun 2026.")
kpi("completed", "Completed Orders", f'=COUNTIF({st},"Completed")', "count", None, None, "Orders with status Completed (= installed)", "Orders fully delivered and installed", "Status = Completed", f"{K['completed']} orders fully closed.")
kpi("pending", "Pending Orders", f'=COUNTIF({st},"Pending")', "count", None, None, "Orders with status Pending", "Open orders still being worked on", "Status = Pending ('In Progress' standardised to Pending)", f"{K['pending']} orders still open.")
kpi("cancelled", "Cancelled Orders", f'=COUNTIF({st},"Cancelled")', "count", None, None, "Orders with status Cancelled", "Lost or withdrawn orders", "Status = Cancelled", f"{K['cancelled']} cancellations - excluded from every delivery and installation KPI.")
kpi("status_check", "Status reconciliation check", f'=IF({ref("completed")}+{ref("pending")}+{ref("cancelled")}={ref("total_orders")},"OK - statuses add up","CHECK - statuses do not add up")', "check", None, None, "Completed + Pending + Cancelled = Total", "Proves no order is lost or double counted", "All orders", "Must always read OK.")
KREF.pop("status_check")
kpi("active_orders", "Active (non-cancelled) Orders", f"={ref('total_orders')}-{ref('cancelled')}", "count", None, None, "Total - Cancelled", "Orders the business still has to fulfil", "Excludes cancelled", f"{K['active_orders']} orders.")
kpi("order_completion_rate", "Order Completion Rate", pct(ref("completed"), ref("active_orders")), "%", f"={ref('completed')}", f"={ref('active_orders')}", "Completed ÷ Active (non-cancelled) orders", "Share of committed orders that are fully closed", "Cancelled orders removed from the denominator; recent orders not yet due still count as open", f"{p(K['order_completion_rate'])}. Lower in June because many June orders are not due yet - read together with Overdue Orders.")
kpi("cancellation_rate", "Cancellation Rate", pct(ref("cancelled"), ref("total_orders")), "%", f"={ref('cancelled')}", f"={ref('total_orders')}", "Cancelled ÷ Total orders", "Lost demand", "All orders", f"{p(K['cancellation_rate'])}; every cancellation has a recorded reason.")
kpi("backlog", "Current Order Backlog (orders)", f"={ref('pending')}", "count", None, None, "Pending orders on the report date", "Open work still to deliver and/or install", "Status = Pending", f"{K['backlog']} open orders; {K['overdue_orders']} of them are overdue (see D).")
kpi("backlog_value", "Backlog Value (INR, simulated)", f'=SUMIFS({RG("Order_Value_INR")},{st},"Pending")', "INR", None, None, "Sum of order value of Pending orders", "Revenue tied up in open orders", "Pending orders with a value", f"₹{K['backlog_value']/1e7:.2f} crore (simulated values).")
section("B. DELIVERY PERFORMANCE")
kpi("deliveries_due", "Total Deliveries Due", f"=SUM({RG('Delivery_Due')})", "count", None, None, "Non-cancelled orders with promised delivery on/before as-of date", "Deliveries the business had committed to make by now", "Excludes cancelled and orders promised after the as-of date (not yet due is NOT overdue)", f"{K['deliveries_due']} deliveries were due.")
kpi("deliveries_due_done", "Due Deliveries Completed", f"=COUNTIFS({RG('Delivery_Due')},1,{RG('Delivered')},1)", "count", None, None, "Due AND delivered", "Commitments met (early, on time or late)", "Delivery_Due = 1", f"{K['deliveries_due_done']} of {K['deliveries_due']}.")
kpi("delivery_completion_rate", "Delivery Completion Rate", pct(ref("deliveries_due_done"), ref("deliveries_due")), "%", f"={ref('deliveries_due_done')}", f"={ref('deliveries_due')}", "Due deliveries completed ÷ Deliveries due", "Did deliveries happen at all (ignores lateness)", "Due orders only", f"{p(K['delivery_completion_rate'])} - almost every due delivery has happened; the issue is lateness, not missed deliveries.")
kpi("deliveries_timed", "Deliveries With Valid Timing", f"=COUNT({RG('Delivery_Delay_Days')})", "count", None, None, "Delivered orders with valid promised and actual dates", "Denominator for on-time %", "Excludes cancelled, undelivered, and rows with an invalid delivery date", f"{K['deliveries_timed']} deliveries.")
kpi("on_time_deliveries", "On-Time Deliveries", f'=COUNTIF({RG("Delivery_Delay_Days")},"<=0")', "count", None, None, "Delay Days <= 0 (early or on the promised date)", "Promises kept", "Timed deliveries", f"{K['on_time_deliveries']}.")
kpi("late_deliveries", "Delayed Delivery Count", f'=COUNTIF({RG("Delivery_Delay_Days")},">0")', "count", None, None, "Delay Days > 0", "Promises missed", "Timed deliveries", f"{K['late_deliveries']} late deliveries.")
kpi("on_time_delivery_rate", "On-Time Delivery Rate", pct(ref("on_time_deliveries"), ref("deliveries_timed")), "%", f"={ref('on_time_deliveries')}", f"={ref('deliveries_timed')}", "On-time deliveries ÷ Deliveries with valid timing", "Reliability of delivery promises", "Pending (undelivered) and cancelled orders are NOT counted as on time or late", f"{p(K['on_time_delivery_rate'])} - roughly 1 in 4 deliveries missed the promise.")
kpi("avg_delivery_delay_late", "Average Delivery Delay (late orders)", f'=IFERROR(AVERAGEIF({RG("Delivery_Delay_Days")},">0"),"n/a")', "days", None, None, "Average Delay Days where Delay Days > 0", "How late a late delivery typically is", "Late deliveries only", f"{K['avg_delivery_delay_late']:.1f} days; a few long component-shortage delays pull this up.")
kpi("avg_delivery_delay_all", "Average Delivery Variance (all timed)", f'=IFERROR(AVERAGE({RG("Delivery_Delay_Days")}),"n/a")', "days", None, None, "Average Delay Days, negative = early", "Overall bias of promises", "Timed deliveries", f"{K['avg_delivery_delay_all']:.1f} days - on average deliveries arrive slightly early, so the problem is concentrated, not universal.")
kpi("max_delivery_delay", "Maximum Delivery Delay", f'=IFERROR(MAX({RG("Delivery_Delay_Days")}),"n/a")', "days", None, None, "Largest Delay Days", "Worst case", "Timed deliveries", f"{K['max_delivery_delay']:.0f} days.")
section("C. INSTALLATION PERFORMANCE")
kpi("installs_due", "Installations Due", f"=SUM({RG('Install_Due')})", "count", None, None, "Non-cancelled orders with promised installation on/before as-of date", "Installations committed by now ('scheduled' = has a promised installation date)", "Excludes cancelled and not-yet-due", f"{K['installs_due']}.")
kpi("installs_due_done", "Due Installations Completed", f"=COUNTIFS({RG('Install_Due')},1,{RG('Installed')},1)", "count", None, None, "Due AND installed", "Installation commitments met", "Install_Due = 1", f"{K['installs_due_done']} of {K['installs_due']}.")
kpi("installation_completion_rate", "Installation Completion Rate", pct(ref("installs_due_done"), ref("installs_due")), "%", f"={ref('installs_due_done')}", f"={ref('installs_due')}", "Due installations completed ÷ Installations due", "Did installations happen when they were due", "Due orders only", f"{p(K['installation_completion_rate'])} - lower than delivery completion ({p(K['delivery_completion_rate'])}).")
kpi("installed", "Completed Installations (all)", f"=SUM({RG('Installed')})", "count", None, None, "Non-cancelled orders installed", "Total installations done", "Includes early installations not yet due", f"{K['installed']}.")
kpi("pending_installations", "Pending Installations", f"=COUNTIFS({RG('Delivered')},1,{RG('Installed')},0)", "count", None, None, "Delivered but not installed", "Machines on customer sites waiting to be commissioned", "Non-cancelled", f"{K['pending_installations']} delivered units await installation.")
kpi("installs_timed", "Installations With Valid Timing", f"=COUNT({RG('Install_Delay_Days')})", "count", None, None, "Installed orders with valid promised and actual dates", "Denominator for on-time install %", "Excludes 2 'installed before delivery' rows and 1 completed order with no install date", f"{K['installs_timed']}.")
kpi("on_time_install_rate", "On-Time Installation Rate", pct(f'COUNTIF({RG("Install_Delay_Days")},"<=0")', ref("installs_timed")), "%", f'=COUNTIF({RG("Install_Delay_Days")},"<=0")', f"={ref('installs_timed')}", "Installations on/before promise ÷ Installations with valid timing", "Reliability of installation promises", "Timed installations", f"{p(K['on_time_install_rate'])} - the weakest stage in the process.")
kpi("avg_install_delay_late", "Average Installation Delay (late)", f'=IFERROR(AVERAGEIF({RG("Install_Delay_Days")},">0"),"n/a")', "days", None, None, "Average Install Delay Days where > 0", "How late a late installation is", "Late installations", f"{K['avg_install_delay_late']:.1f} days.")
kpi("avg_install_duration", "Average Installation Duration", f'=IFERROR(AVERAGE({RG("Install_Lead_Days")}),"n/a")', "days", None, None, "Average of (Installation date - Delivery date)", "Time a delivered machine waits before it is working", "Valid delivery and installation dates", f"{K['avg_install_duration']:.1f} days from delivery to installation.")
section("D. OPERATIONAL EFFICIENCY & PENDING WORK")
kpi("avg_production", "Average Production Time", f'=IFERROR(AVERAGE({RG("Production_Days")}),"n/a")', "days", None, None, "Average (Dispatch - Order date)", "Factory / processing time", "Valid dispatch dates, non-cancelled", f"{K['avg_production']:.1f} days - the largest part of the cycle.")
kpi("avg_transit", "Average Transit Time", f'=IFERROR(AVERAGE({RG("Transit_Days")}),"n/a")', "days", None, None, "Average (Delivery - Dispatch)", "Logistics time", "Valid dispatch and delivery dates", f"{K['avg_transit']:.1f} days.")
kpi("avg_order_to_delivery", "Average Order Fulfilment Time", f'=IFERROR(AVERAGE({RG("Order_To_Delivery_Days")}),"n/a")', "days", None, None, "Average (Delivery - Order date)", "How long customers wait for the machine", "Valid delivery dates", f"{K['avg_order_to_delivery']:.1f} days.")
kpi("avg_order_cycle", "Average Order Cycle Time", f'=IFERROR(AVERAGE({RG("Order_Cycle_Days")}),"n/a")', "days", None, None, "Average (Installation - Order date) for completed orders", "End-to-end time to a working system", "Completed orders with valid install date", f"{K['avg_order_cycle']:.1f} days end to end.")
kpi("overdue_orders", "Overdue Orders", f"=SUM({RG('Overdue')})", "count", None, None, "Pending orders past the promise for their current stage", "Open orders already breaking a promise", "Pending only; orders whose promise is after the as-of date are NOT overdue", f"{K['overdue_orders']} overdue orders.")
kpi("overdue_delivery_stage", "  ...overdue at delivery stage", f'=COUNTIF({RG("Overdue_Stage")},"Delivery")', "count", None, None, "Overdue and not yet delivered", "Factory/logistics backlog", "Overdue orders", f"{K['overdue_delivery_stage']}.")
kpi("overdue_install_stage", "  ...overdue at installation stage", f'=COUNTIF({RG("Overdue_Stage")},"Installation")', "count", None, None, "Overdue and delivered, not installed", "Site/installation backlog", "Overdue orders", f"{K['overdue_install_stage']} - most overdue work is waiting for installation.")
kpi("avg_days_overdue", "Average Days Overdue", f'=IFERROR(AVERAGE({RG("Days_Overdue")}),"n/a")', "days", None, None, "Average (As-of date - promise date) for overdue orders", "How stale the overdue backlog is", "Overdue orders", f"{K['avg_days_overdue']:.0f} days on average; oldest {K['max_days_overdue']:.0f} days.")
section("E. ORDER VALUE (simulated, illustrative)")
kpi("total_order_value", "Total Order Value", f'=SUMIFS({RG("Order_Value_INR")},{st},"<>Cancelled")', "INR", None, None, "Sum of order value, non-cancelled", "Size of business handled", "3 orders with missing value are excluded (see next row)", f"₹{K['total_order_value']/1e7:.2f} crore (simulated - not real revenue).")
kpi("orders_missing_value", "Orders Missing a Value", f'=SUMPRODUCT(({RG("Order_ID")}<>"")*({RG("Order_Value_INR")}=""))', "count", None, None, "Rows with blank order value", "Completeness of value data", "All orders", f"{K['orders_missing_value']} rows - value KPIs are slightly understated.")
kpi("completed_value", "Completed Order Value", f'=SUMIFS({RG("Order_Value_INR")},{st},"Completed")', "INR", None, None, "Sum of order value, Completed", "Value fully delivered and installed", "Completed orders", f"₹{K['completed_value']/1e7:.2f} crore.")
section("F. DATA QUALITY")
kpi("dq_rows", "Rows With a Data-Quality Flag", f'=COUNTIFS({RG("DQ_Flag")},"<>OK",{RG("DQ_Flag")},"<>")', "count", None, None, "DQ_Flag not OK", "Rows needing confirmation with the data owner", "All orders", f"{K['dq_rows']} rows; each is excluded only from the metric it would distort.")
kpi("late_del_noreason", "Late Deliveries Without a Reason", f'=COUNTIFS({RG("Delivery_Delay_Days")},">0",{RG("Delivery_Delay_Reason")},"")', "count", None, None, "Late deliveries with blank reason", "Root cause cannot be analysed for these", "Late deliveries", f"{K['late_del_noreason']} of {K['late_deliveries']} - make the reason mandatory.")
kpi("late_inst_noreason", "Late Installations Without a Reason", f'=COUNTIFS({RG("Install_Delay_Days")},">0",{RG("Installation_Delay_Reason")},"")', "count", None, None, "Late installations with blank reason", "Root cause cannot be analysed for these", "Late installations", f"{K['late_inst_noreason']} of {K['late_installs']} - make the reason mandatory.")
for j, w in enumerate([5, 30, 14, 7, 12, 12, 34, 32, 36, 46], start=1): ws.column_dimensions[CL(j)].width = w
ws.freeze_panes = "C7"
KPI_LAST = row

# ======================= PIVOT ANALYSIS =======================
ws = S["PIVOT ANALYSIS"]
title(ws, "PIVOT ANALYSIS - formula-built summary tables (whole dataset, not filtered)", SYN)
ws["A3"] = ("These tables are built with COUNTIFS / AVERAGEIFS formulas so they always match CLEANED DATA. "
            "A real PivotTable + slicer can be added in 2 minutes - steps are at the bottom of this sheet.")
ws["A3"].font = font(9, i=True)
r0 = 5
def table(title_txt, dim_col, labels, metrics, label_fmt=None):
    """metrics: list of (header, formula_template_with_{c}, fmt) where {c} is the criteria cell"""
    global r0
    ws.cell(r0, 1, title_txt).font = font(11, True, TEAL); r0 += 1
    header_row(ws, r0, [dim_col] + [m[0] for m in metrics]); r0 += 1
    first = r0
    for lab in labels:
        c = ws.cell(r0, 1, lab); c.font = font(9, True); c.border = BOX
        if label_fmt: c.number_format = label_fmt
        for j, (h, f, fmt) in enumerate(metrics, start=2):
            cell = ws.cell(r0, j, f.replace("{c}", f"$A{r0}")); cell.number_format = fmt; cell.font = font(9); cell.border = BOX
        r0 += 1
    # total row
    ws.cell(r0, 1, "Total / overall").font = font(9, True)
    for j, (h, f, fmt) in enumerate(metrics, start=2):
        tf = f.replace(",{DIM},{c}", "").replace("{DIM},{c},", "")
        cell = ws.cell(r0, j, tf); cell.number_format = fmt; cell.font = font(9, True); cell.fill = fill(SILVER); cell.border = BOX
    r0 += 2
    return first

def dimmetrics(dim):
    D = RG(dim)
    m = [
     ("Orders", f"=COUNTIFS({{DIM}},{{c}})", INT),
     ("Completed", f'=COUNTIFS({{DIM}},{{c}},{st},"Completed")', INT),
     ("Pending", f'=COUNTIFS({{DIM}},{{c}},{st},"Pending")', INT),
     ("Cancelled", f'=COUNTIFS({{DIM}},{{c}},{st},"Cancelled")', INT),
     ("Order completion %", f'=IFERROR(COUNTIFS({{DIM}},{{c}},{st},"Completed")/(COUNTIFS({{DIM}},{{c}})-COUNTIFS({{DIM}},{{c}},{st},"Cancelled")),"")', PCT),
     ("Timed deliveries", f'=COUNTIFS({{DIM}},{{c}},{RG("Delivery_Delay_Days")},">-9999")', INT),
     ("On-time delivery %", f'=IFERROR(COUNTIFS({{DIM}},{{c}},{RG("Delivery_Delay_Days")},"<=0")/COUNTIFS({{DIM}},{{c}},{RG("Delivery_Delay_Days")},">-9999"),"")', PCT),
     ("Late deliveries", f'=COUNTIFS({{DIM}},{{c}},{RG("Delivery_Delay_Days")},">0")', INT),
     ("Avg delay when late (d)", f'=IFERROR(AVERAGEIFS({RG("Delivery_Delay_Days")},{{DIM}},{{c}},{RG("Delivery_Delay_Days")},">0"),"")', DEC),
     ("On-time install %", f'=IFERROR(COUNTIFS({{DIM}},{{c}},{RG("Install_Delay_Days")},"<=0")/COUNTIFS({{DIM}},{{c}},{RG("Install_Delay_Days")},">-9999"),"")', PCT),
     ("Avg production (d)", f'=IFERROR(AVERAGEIFS({RG("Production_Days")},{{DIM}},{{c}}),"")', DEC),
     ("Avg cycle (d)", f'=IFERROR(AVERAGEIFS({RG("Order_Cycle_Days")},{{DIM}},{{c}}),"")', DEC),
     ("Overdue", f'=COUNTIFS({{DIM}},{{c}},{RG("Overdue")},1)', INT),
     ("Order value ₹ (non-cancelled)", f'=SUMIFS({RG("Order_Value_INR")},{{DIM}},{{c}},{st},"<>Cancelled")', INR),
    ]
    out = []
    for h, f, fmt in m:
        f = f.replace("{DIM}", "{DIM}")  # keep marker for total-row stripping
        out.append((h, f, fmt))
    return out, D
def build(title_txt, dim, labels, label_fmt=None):
    ms, D = dimmetrics(dim)
    ms2 = [(h, f.replace("{DIM}", D) if False else f, fmt) for h, f, fmt in ms]
    # total row must drop the dimension criterion; do substitution after total computed
    global r0
    ws.cell(r0, 1, title_txt).font = font(11, True, TEAL); r0 += 1
    header_row(ws, r0, [dim.replace("_", " ")] + [m[0] for m in ms]); r0 += 1
    for lab in labels:
        c = ws.cell(r0, 1, lab); c.font = font(9, True); c.border = BOX
        if label_fmt: c.number_format = label_fmt
        for j, (h, f, fmt) in enumerate(ms, start=2):
            cell = ws.cell(r0, j, f.replace("{DIM}", D).replace("{c}", f"$A{r0}")); cell.number_format = fmt; cell.font = font(9); cell.border = BOX
        r0 += 1
    ws.cell(r0, 1, "Overall").font = font(9, True); ws.cell(r0, 1).fill = fill(SILVER)
    for j, (h, f, fmt) in enumerate(ms, start=2):
        tf = f.replace("COUNTIFS({DIM},{c})", f"COUNTA({RG('Order_ID')})").replace("{DIM},{c},", "").replace(",{DIM},{c}", "")
        tf = tf.replace("AVERAGEIFS(" + RG("Production_Days") + ")", "AVERAGE(" + RG("Production_Days") + ")") \
               .replace("AVERAGEIFS(" + RG("Order_Cycle_Days") + ")", "AVERAGE(" + RG("Order_Cycle_Days") + ")")
        cell = ws.cell(r0, j, tf); cell.number_format = fmt; cell.font = font(9, True); cell.fill = fill(SILVER); cell.border = BOX
    r0 += 2
PIV = {}
PIV["cat"] = r0 + 2; build("1. Performance by PRODUCT CATEGORY", "Product_Category", CATS)
PIV["month"] = r0 + 2; build("2. Performance by ORDER MONTH", "Order_Month", [datetime(2026, m, 1) for m in range(1, 7)], "mmm-yyyy")
PIV["ctype"] = r0 + 2; build("3. Performance by CUSTOMER TYPE", "Customer_Type", CTYPES)
PIV["region"] = r0 + 2; build("4. Performance by REGION", "Region", REGIONS)
# 5 status x stage
ws.cell(r0, 1, "5. ORDER STATUS x CURRENT STAGE (where open work is sitting)").font = font(11, True, TEAL); r0 += 1
stages = ["In Production", "In Transit", "Awaiting Installation", "Completed", "Cancelled"]
header_row(ws, r0, ["Current stage", "Orders", "of which overdue", "Avg days overdue"]); r0 += 1
PIV["stage"] = r0
for sname in stages:
    ws.cell(r0, 1, sname).font = font(9, True)
    ws.cell(r0, 2, f"=COUNTIF({RG('Current_Stage')},$A{r0})").number_format = INT
    ws.cell(r0, 3, f"=COUNTIFS({RG('Current_Stage')},$A{r0},{RG('Overdue')},1)").number_format = INT
    ws.cell(r0, 4, f'=IFERROR(AVERAGEIFS({RG("Days_Overdue")},{RG("Current_Stage")},$A{r0}),"")').number_format = DEC
    for j in range(1, 5): ws.cell(r0, j).border = BOX; ws.cell(r0, j).font = font(9, j == 1)
    r0 += 1
r0 += 1
# 6 delay reasons
ws.cell(r0, 1, "6. DELAY REASONS (late records only)").font = font(11, True, TEAL); r0 += 1
header_row(ws, r0, ["Delivery delay reason", "Late deliveries", "Avg delay (d)", "", "Installation delay reason", "Late installations", "Avg delay (d)"]); r0 += 1
dres = ["Production backlog", "Component shortage", "Transport scheduling", "Dispatch scheduling", ""]
ires = ["Customer site not ready", "Knock-on from late delivery", "Technician availability", "Installation scheduling", ""]
PIV["reasons"] = r0
for a, b in zip(dres, ires):
    ws.cell(r0, 1, a or "(not recorded)").font = font(9, True)
    crit_a = f'"{a}"' if a else '""'
    ws.cell(r0, 2, f'=COUNTIFS({RG("Delivery_Delay_Days")},">0",{RG("Delivery_Delay_Reason")},{crit_a})').number_format = INT
    ws.cell(r0, 3, f'=IFERROR(AVERAGEIFS({RG("Delivery_Delay_Days")},{RG("Delivery_Delay_Days")},">0",{RG("Delivery_Delay_Reason")},{crit_a}),"")').number_format = DEC
    ws.cell(r0, 5, b or "(not recorded)").font = font(9, True)
    crit_b = f'"{b}"' if b else '""'
    ws.cell(r0, 6, f'=COUNTIFS({RG("Install_Delay_Days")},">0",{RG("Installation_Delay_Reason")},{crit_b})').number_format = INT
    ws.cell(r0, 7, f'=IFERROR(AVERAGEIFS({RG("Install_Delay_Days")},{RG("Install_Delay_Days")},">0",{RG("Installation_Delay_Reason")},{crit_b}),"")').number_format = DEC
    for j in (1, 2, 3, 5, 6, 7): ws.cell(r0, j).border = BOX
    r0 += 1
ws.cell(r0, 1, "Note: 'Dispatch scheduling' and 'Installation scheduling' are valid reason codes that were not used in this period.").font = font(8, i=True); r0 += 2
# 7 stage comparison
ws.cell(r0, 1, "7. STAGE COMPARISON - which stage causes the most delay?").font = font(11, True, TEAL); r0 += 1
header_row(ws, r0, ["Stage", "Avg duration (d)", "Records timed", "Late vs promise", "Late share", "Overdue open orders"]); r0 += 1
PIV["stagecmp"] = r0
rows_sc = [("Production (order -> dispatch)", "Production_Days", None, None, None),
           ("Transit (dispatch -> delivery)", "Transit_Days", None, None, None),
           ("Delivery promise (customer-facing)", "Order_To_Delivery_Days", "Delivery_Delay_Days", None, "Delivery"),
           ("Installation (delivery -> installed)", "Install_Lead_Days", "Install_Delay_Days", None, "Installation")]
for lab, durc, latec, _, ov in rows_sc:
    ws.cell(r0, 1, lab).font = font(9, True)
    ws.cell(r0, 2, f'=IFERROR(AVERAGE({RG(durc)}),"")').number_format = DEC
    ws.cell(r0, 3, f"=COUNT({RG(durc)})").number_format = INT
    if latec:
        ws.cell(r0, 4, f'=COUNTIF({RG(latec)},">0")').number_format = INT
        ws.cell(r0, 5, f'=IFERROR(COUNTIF({RG(latec)},">0")/COUNT({RG(latec)}),"")').number_format = PCT
        ws.cell(r0, 6, f'=COUNTIF({RG("Overdue_Stage")},"{ov}")').number_format = INT
    else:
        for j in (4, 5, 6): ws.cell(r0, j, "n/a (no stage promise date)").font = font(8, i=True, c=MID)
    for j in range(1, 7): ws.cell(r0, j).border = BOX
    r0 += 1
r0 += 1
ws.cell(r0, 1, "HOW TO ADD A REAL PIVOTTABLE + SLICER (Excel desktop, about 2 minutes)").font = font(11, True, TEAL); r0 += 1
for s_ in ["1. Go to CLEANED DATA, click any cell inside the data, press Ctrl+A.",
           "2. Insert > PivotTable > New Worksheet > OK.",
           "3. Drag Product_Category to Rows, Order_Status to Columns, Order_ID to Values (it shows 'Count of Order_ID').",
           "4. Click inside the PivotTable > PivotTable Analyze > Insert Slicer > tick Region and Customer_Type > OK.",
           "5. Click slicer buttons to filter. The counts must match table 1 above - that is your reconciliation check."]:
    ws.cell(r0, 1, s_).font = font(9); r0 += 1
ws.column_dimensions["A"].width = 34
for j in range(2, 16): ws.column_dimensions[CL(j)].width = 13
ws.column_dimensions["E"].width = 26

# ======================= OPERATIONS DASHBOARD =======================
ws = S["OPERATIONS DASHBOARD"]
for j in range(1, 40): ws.column_dimensions[CL(j)].width = 10.5
ws.column_dimensions["A"].width = 2
ws.merge_cells("B1:Q1"); ws["B1"] = "ORDER FULFILMENT & INSTALLATION PERFORMANCE DASHBOARD"
ws["B1"].font = font(18, True, "FFFFFF"); ws["B1"].fill = fill(CHAR); ws["B1"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 34
ws.merge_cells("B2:Q2")
ws["B2"] = '="Neo Heat Pump - Sai Heating and Cooling Solutions  |  Case study  |  Orders Jan-Jun 2026  |  As of "&TEXT(AsOf,"dd-mmm-yyyy")&"  |  SIMULATED DATA - not real company performance"'
ws["B2"].font = font(9, True, "FFFFFF"); ws["B2"].fill = fill(TEAL); ws["B2"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[2].height = 20
# filters
for col, lab, nm, lst in [("B", "Category", "fCat", "$A$2:$A$7"), ("G", "Customer", "fCust", "$B$2:$B$5"), ("L", "Region", "fReg", "$C$2:$C$5")]:
    ws[f"{col}4"] = lab; ws[f"{col}4"].font = font(9, True, MID)
    c2 = CL(ws[f"{col}4"].column + 1); c3 = CL(ws[f"{col}4"].column + 3)
    ws.merge_cells(f"{c2}4:{c3}4"); ws[f"{c2}4"] = "All"
    ws[f"{c2}4"].font = font(10, True, "0000FF"); ws[f"{c2}4"].fill = fill("FFF59D"); ws[f"{c2}4"].border = BOX
    ws[f"{c2}4"].alignment = CENTER
    dvx = DataValidation(type="list", formula1=f"=LISTS!{lst.replace('$A', '$A').replace('$B', '$B').replace('$C', '$C')}", allow_blank=False)
    ws.add_data_validation(dvx); dvx.add(f"{c2}4")
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'OPERATIONS DASHBOARD'!${c2}$4")
ws["B5"] = "Pick a value in the yellow boxes to filter every card, chart and list on this sheet. Choose 'All' to reset."
ws["B5"].font = font(8, i=True, c=MID)
IF_ = RG("In_Dashboard_Filter")
def cf(*crit):  # filtered count
    return f"COUNTIFS({IF_},1" + "".join(f",{RG(a)},{b}" for a, b in crit) + ")"
cards = [
 ("TOTAL ORDERS", f"={cf()}", INT, f'="Jan-Jun 2026"'),
 ("ORDER COMPLETION", f'=IFERROR({cf(("Order_Status", chr(34)+"Completed"+chr(34)))}/({cf()}-{cf(("Order_Status", chr(34)+"Cancelled"+chr(34)))}),"n/a")', PCT,
  f'={cf(("Order_Status", chr(34)+"Completed"+chr(34)))}&" completed · "&{cf(("Order_Status", chr(34)+"Cancelled"+chr(34)))}&" cancelled"'),
 ("OPEN BACKLOG", f'={cf(("Order_Status", chr(34)+"Pending"+chr(34)))}', INT, f'="₹"&TEXT(SUMIFS({RG("Order_Value_INR")},{IF_},1,{st},"Pending")/100000,"#,##0")&" lakh open value"'),
 ("ON-TIME DELIVERY", f'=IFERROR({cf(("Delivery_Delay_Days", chr(34)+"<=0"+chr(34)))}/{cf(("Delivery_Delay_Days", chr(34)+">-9999"+chr(34)))},"n/a")', PCT,
  f'={cf(("Delivery_Delay_Days", chr(34)+">0"+chr(34)))}&" late of "&{cf(("Delivery_Delay_Days", chr(34)+">-9999"+chr(34)))}&" delivered"'),
 ("INSTALL COMPLETION", f'=IFERROR({cf(("Install_Due", "1"), ("Installed", "1"))}/{cf(("Install_Due", "1"))},"n/a")', PCT,
  f'={cf(("Install_Due", "1"), ("Installed", "1"))}&" of "&{cf(("Install_Due", "1"))}&" due installs done"'),
 ("ON-TIME INSTALL", f'=IFERROR({cf(("Install_Delay_Days", chr(34)+"<=0"+chr(34)))}/{cf(("Install_Delay_Days", chr(34)+">-9999"+chr(34)))},"n/a")', PCT,
  f'="avg "&TEXT(IFERROR(AVERAGEIFS({RG("Install_Lead_Days")},{IF_},1),0),"0.0")&" days delivery→install"'),
 ("OVERDUE ORDERS", f'={cf(("Overdue", "1"))}', INT, f'={cf(("Overdue_Stage", chr(34)+"Installation"+chr(34)))}&" stuck at installation"'),
 ("ORDER VALUE", f'=SUMIFS({RG("Order_Value_INR")},{IF_},1,{st},"<>Cancelled")/10000000', '"₹"0.00" Cr"', f'="simulated · excl. cancelled"'),
]
for n, (lab, fml, fmt, sub) in enumerate(cards):
    c1 = 2 + n * 2; a, b = CL(c1), CL(c1 + 1)
    for rr in (7, 8, 9):
        ws.merge_cells(f"{a}{rr}:{b}{rr}")
        for cc in (c1, c1 + 1): ws.cell(rr, cc).fill = fill(SILVER)
    ws[f"{a}7"] = lab; ws[f"{a}7"].font = font(8, True, MID); ws[f"{a}7"].alignment = CENTER
    ws[f"{a}8"] = fml; ws[f"{a}8"].number_format = fmt; ws[f"{a}8"].alignment = CENTER
    ws[f"{a}8"].font = font(18, True, RED if lab == "OVERDUE ORDERS" else (TEAL if fmt == PCT else CHAR))
    ws[f"{a}9"] = sub; ws[f"{a}9"].font = font(8, c=MID); ws[f"{a}9"].alignment = CENTER
    ws.cell(6, c1).fill = fill(TEAL); ws.cell(6, c1 + 1).fill = fill(TEAL)
ws.row_dimensions[6].height = 4; ws.row_dimensions[8].height = 30; ws.row_dimensions[9].height = 24
# ---- chart data (column T onward) ----
CD = 20  # column T
ws.cell(11, CD, "CHART DATA - formulas, follow the filters (do not edit)").font = font(9, True, MID)
def put(r, c, v, fmt=None, b=False):
    cell = ws.cell(r, c, v); cell.font = font(8, b, MID if not b else CHAR)
    if fmt: cell.number_format = fmt
    return cell
# A status
r = 13; put(r - 1, CD, "Order status", b=True); put(r - 1, CD + 1, "Orders", b=True)
for i, sname in enumerate(["Completed", "Pending", "Cancelled"]):
    put(r + i, CD, sname); put(r + i, CD + 1, f'={cf(("Order_Status", f"$T{r+i}"))}', INT)
STATUS_RNG = (r, r + 2)
# B month
r = 18; put(r - 1, CD, "Month", b=True)
for j, sname in enumerate(["Completed", "Pending", "Cancelled"]): put(r - 1, CD + 1 + j, sname, b=True)
for i in range(6):
    put(r + i, CD, datetime(2026, i + 1, 1), "mmm")
    for j, sname in enumerate(["Completed", "Pending", "Cancelled"]):
        put(r + i, CD + 1 + j, f'={cf(("Order_Month", f"$T{r+i}"), ("Order_Status", f"{CL(CD+1+j)}${r-1}"))}', INT)
MONTH_RNG = (r, r + 5)
# C delivery by category
r = 26; put(r - 1, CD, "Category", b=True); put(r - 1, CD + 1, "On time / early", b=True); put(r - 1, CD + 2, "Late", b=True)
SHORT = {"Heat Pump Water Heater": "Water Heater", "Commercial Heat Pump": "Commercial HP", "Swimming Pool Heat Pump": "Pool HP",
         "Heat Pump Dryer": "Dryer", "Custom Requirement": "Custom"}
for i, cname in enumerate(CATS):
    put(r + i, CD, SHORT[cname]); put(r + i, CD + 4, cname)
    put(r + i, CD + 1, f'={cf(("Product_Category", f"$X{r+i}"), ("Delivery_Delay_Days", chr(34)+"<=0"+chr(34)))}', INT)
    put(r + i, CD + 2, f'={cf(("Product_Category", f"$X{r+i}"), ("Delivery_Delay_Days", chr(34)+">0"+chr(34)))}', INT)
DELCAT_RNG = (r, r + 4)
# D installation status
r = 33; put(r - 1, CD, "Installation status", b=True); put(r - 1, CD + 1, "Orders", b=True)
inst_rows = [("Installed on time", f'={cf(("Install_Delay_Days", chr(34)+"<=0"+chr(34)))}'),
             ("Installed late", f'={cf(("Install_Delay_Days", chr(34)+">0"+chr(34)))}'),
             ("Installed - date issue", f'={cf(("Installed", "1"))}-{cf(("Install_Delay_Days", chr(34)+">-9999"+chr(34)))}'),
             ("Awaiting install - overdue", f'={cf(("Overdue_Stage", chr(34)+"Installation"+chr(34)))}'),
             ("Awaiting install - on plan", f'={cf(("Current_Stage", chr(34)+"Awaiting Installation"+chr(34)))}-{cf(("Overdue_Stage", chr(34)+"Installation"+chr(34)))}'),
             ("Not yet delivered", f'={cf(("Current_Stage", chr(34)+"In Production"+chr(34)))}+{cf(("Current_Stage", chr(34)+"In Transit"+chr(34)))}')]
for i, (lab, f_) in enumerate(inst_rows):
    put(r + i, CD, lab); put(r + i, CD + 1, f_, INT)
INST_RNG = (r, r + 5)
# E category performance %
r = 42; put(r - 1, CD, "Category", b=True); put(r - 1, CD + 1, "On-time delivery %", b=True); put(r - 1, CD + 2, "On-time install %", b=True)
for i, cname in enumerate(CATS):
    put(r + i, CD, SHORT[cname]); put(r + i, CD + 4, cname)
    put(r + i, CD + 1, f'=IFERROR({cf(("Product_Category", f"$X{r+i}"), ("Delivery_Delay_Days", chr(34)+"<=0"+chr(34)))}/{cf(("Product_Category", f"$X{r+i}"), ("Delivery_Delay_Days", chr(34)+">-9999"+chr(34)))},0)', PCT)
    put(r + i, CD + 2, f'=IFERROR({cf(("Product_Category", f"$X{r+i}"), ("Install_Delay_Days", chr(34)+"<=0"+chr(34)))}/{cf(("Product_Category", f"$X{r+i}"), ("Install_Delay_Days", chr(34)+">-9999"+chr(34)))},0)', PCT)
CATPERF_RNG = (r, r + 4)
# F stage durations
r = 50; put(r - 1, CD, "Stage", b=True); put(r - 1, CD + 1, "Avg days", b=True)
for i, (lab, c_) in enumerate([("Production", "Production_Days"), ("Transit", "Transit_Days"), ("Delivery → install", "Install_Lead_Days")]):
    put(r + i, CD, lab); put(r + i, CD + 1, f'=IFERROR(AVERAGEIFS({RG(c_)},{IF_},1),0)', DEC)
STAGE_RNG = (r, r + 2)

def style_chart(ch, ttl, h=7.2, w=15.6):
    ch.title = ttl; ch.height = h; ch.width = w; ch.legend.position = "b"
    return ch
def colorize(ch, colors):
    for s_, c_ in zip(ch.series, colors):
        s_.graphicalProperties.solidFill = c_; s_.graphicalProperties.line.solidFill = c_
# chart 1 status doughnut -> use bar for clarity
ch = BarChart(); ch.type = "bar"; ch.style = 10
ch.add_data(Reference(ws, min_col=CD + 1, min_row=STATUS_RNG[0] - 1, max_row=STATUS_RNG[1]), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=CD, min_row=STATUS_RNG[0], max_row=STATUS_RNG[1]))
style_chart(ch, "Order status distribution"); ch.legend = None; colorize(ch, [TEAL])
ch.dataLabels = DLL(); ch.y_axis.majorGridlines = None
ch.x_axis.delete = False; ch.y_axis.delete = False
ws.add_chart(ch, "B11")
ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
ch.add_data(Reference(ws, min_col=CD + 1, max_col=CD + 3, min_row=MONTH_RNG[0] - 1, max_row=MONTH_RNG[1]), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=CD, min_row=MONTH_RNG[0], max_row=MONTH_RNG[1]))
style_chart(ch, "Monthly orders by current status"); colorize(ch, [TEAL, AMBER, MID])
ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.number_format = "mmm"
ws.add_chart(ch, "J11")
ch = BarChart(); ch.type = "bar"; ch.grouping = "stacked"; ch.overlap = 100
ch.add_data(Reference(ws, min_col=CD + 1, max_col=CD + 2, min_row=DELCAT_RNG[0] - 1, max_row=DELCAT_RNG[1]), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=CD, min_row=DELCAT_RNG[0], max_row=DELCAT_RNG[1]))
style_chart(ch, "On-time vs delayed deliveries by product"); colorize(ch, [TEAL, RED])
ch.x_axis.delete = False; ch.y_axis.delete = False
ws.add_chart(ch, "B26")
ch = BarChart(); ch.type = "bar"
ch.add_data(Reference(ws, min_col=CD + 1, min_row=INST_RNG[0] - 1, max_row=INST_RNG[1]), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=CD, min_row=INST_RNG[0], max_row=INST_RNG[1]))
style_chart(ch, "Installation status (non-cancelled orders)"); ch.legend = None; colorize(ch, [TEAL])
ch.dataLabels = DLL()
ch.x_axis.delete = False; ch.y_axis.delete = False
ws.add_chart(ch, "J26")
ch = BarChart(); ch.type = "col"
ch.add_data(Reference(ws, min_col=CD + 1, max_col=CD + 2, min_row=CATPERF_RNG[0] - 1, max_row=CATPERF_RNG[1]), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=CD, min_row=CATPERF_RNG[0], max_row=CATPERF_RNG[1]))
style_chart(ch, "Product-category performance (on-time %)"); colorize(ch, [TEAL, CHAR])
ch.y_axis.number_format = "0%"; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1
ch.x_axis.delete = False; ch.y_axis.delete = False
ws.add_chart(ch, "B41")
ch = BarChart(); ch.type = "col"
ch.add_data(Reference(ws, min_col=CD + 1, min_row=STAGE_RNG[0] - 1, max_row=STAGE_RNG[1]), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=CD, min_row=STAGE_RNG[0], max_row=STAGE_RNG[1]))
style_chart(ch, "Operational delay analysis: average days per stage"); ch.legend = None; colorize(ch, [AMBER])
ch.dataLabels = DLL()
ch.x_axis.delete = False; ch.y_axis.delete = False
ws.add_chart(ch, "J41")
# ---- outstanding orders table ----
r = 57
ws.merge_cells(f"B{r}:Q{r}"); ws[f"B{r}"] = "OUTSTANDING ORDERS - most overdue first (follows filters)"
ws[f"B{r}"].font = font(11, True, "FFFFFF"); ws[f"B{r}"].fill = fill(CHAR); ws[f"B{r}"].alignment = Alignment(indent=1)
heads = [("Order ID", "B", "C", "Order_ID"), ("Product", "D", "F", "Product_Category"), ("Customer type", "G", "H", "Customer_Type"),
         ("Region", "I", "J", "Region"), ("Stuck at", "K", "L", "Overdue_Stage"), ("Days overdue", "M", "M", "Days_Overdue"),
         ("Recorded reason", "N", "Q", None)]
for h, a, b, _ in heads:
    if a != b: ws.merge_cells(f"{a}{r+1}:{b}{r+1}")
    ws[f"{a}{r+1}"] = h; ws[f"{a}{r+1}"].font = font(9, True, "FFFFFF"); ws[f"{a}{r+1}"].fill = fill(TEAL); ws[f"{a}{r+1}"].alignment = CENTER
KEY = RG("Overdue_Rank_Key")
for k in range(1, 11):
    rr = r + 1 + k
    ws[f"S{rr}"] = f'=IF(LARGE({KEY},{k})>0,MATCH(LARGE({KEY},{k}),{KEY},0),"")'; ws[f"S{rr}"].font = font(7, c="BBBBBB")
    for h, a, b, c_ in heads:
        if a != b: ws.merge_cells(f"{a}{rr}:{b}{rr}")
        if c_:
            f_ = f'=IF($S{rr}="","",INDEX({RG(c_)},$S{rr}))'
        else:
            f_ = (f'=IF($S{rr}="","",IF(INDEX({RG("Overdue_Stage")},$S{rr})="Installation",'
                  f'IF(INDEX({RG("Installation_Delay_Reason")},$S{rr})="","(not recorded)",INDEX({RG("Installation_Delay_Reason")},$S{rr})),'
                  f'IF(INDEX({RG("Delivery_Delay_Reason")},$S{rr})="","(not recorded)",INDEX({RG("Delivery_Delay_Reason")},$S{rr}))))')
        cell = ws[f"{a}{rr}"]; cell.value = f_; cell.font = font(9, c_ == "Order_ID", RED if c_ == "Days_Overdue" else CHAR)
        cell.alignment = Alignment(horizontal="center" if c_ in ("Days_Overdue", "Order_ID") else "left", vertical="center")
        cell.number_format = "0"
        for cc in range(ws[f"{a}{rr}"].column, ws[f"{b}{rr}"].column + 1):
            ws.cell(rr, cc).border = Border(bottom=thin)
ws.cell(r + 12, 2, '=IF(SUM(' + RG("Overdue_Rank_Key") + ')=0,"No overdue orders for this filter selection.","Showing up to 10 overdue orders. Full list: filter CLEANED DATA column Overdue = 1.")').font = font(8, i=True, c=MID)
ws.cell(r + 14, 2, "How to read: cards = headline KPIs; charts = where the problems are; table = what to chase today. Definitions: KPI ANALYSIS sheet.").font = font(8, i=True, c=MID)
ws.freeze_panes = "A6"
ws.print_area = "A1:Q72"
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ======================= BUSINESS RECOMMENDATIONS =======================
ws = S["BUSINESS RECOMMENDATIONS"]
title(ws, "BUSINESS FINDINGS & RECOMMENDATIONS", "Findings describe SIMULATED data. Recommendations are PROPOSALS - none has been implemented, and no improvement is claimed.")
cat_rows = {c: PIV["cat"] + i for i, c in enumerate(CATS)}
def pv(cat, col):  # pivot cell reference
    letter = {"ontime": "H", "late": "I", "orders": "B", "prod": "L", "cycle": "M"}[col]
    return f"'PIVOT ANALYSIS'!${letter}${cat_rows[cat]}"
reg_rows = {c: PIV["region"] + i for i, c in enumerate(REGIONS)}
big = R["big_vs_rest"]
FIND = [
 ("Delivery promises are missed on roughly 1 in 4 deliveries.",
  f'="On-time delivery "&TEXT({KREF["on_time_delivery_rate"]},"0.0%")&" ("&{KREF["on_time_deliveries"]}&" of "&{KREF["deliveries_timed"]}&"); "&{KREF["late_deliveries"]}&" late, avg "&TEXT({KREF["avg_delivery_delay_late"]},"0.0")&" days late"',
  "Observed (calculated)", "Promised lead times may not reflect real production time for some products (hypothesis).",
  "Missed promises create follow-up calls and erode trust (impact not measured in this data)."),
 ("Late deliveries are concentrated in Commercial Heat Pumps and Custom Requirements.",
  f'="Commercial HP "&TEXT({pv("Commercial Heat Pump","ontime")},"0%")&" and Custom "&TEXT({pv("Custom Requirement","ontime")},"0%")&" on time vs Water Heater "&TEXT({pv("Heat Pump Water Heater","ontime")},"0%")&"; these two categories = "&({pv("Commercial Heat Pump","late")}+{pv("Custom Requirement","late")})&" of "&{KREF["late_deliveries"]}&" late deliveries"',
  "Observed (calculated)", "Longer, more variable production (avg production days are highest for these categories) and component availability (hypothesis; built into the simulation).",
  "Category-specific promise dates could protect on-time performance where it is weakest."),
 ("Installation is the weakest stage of the process.",
  f'="On-time installation "&TEXT({KREF["on_time_install_rate"]},"0.0%")&" vs on-time delivery "&TEXT({KREF["on_time_delivery_rate"]},"0.0%")&"; installation completion "&TEXT({KREF["installation_completion_rate"]},"0.0%")&" vs delivery completion "&TEXT({KREF["delivery_completion_rate"]},"0.0%")',
  "Observed (calculated)", "Part of this is knock-on from late deliveries; part is installation-stage waiting (see next row).",
  "A machine delivered but not installed is not yet useful to the customer."),
 ("Many late installations happened even though delivery was on time.",
  f'=COUNTIFS({RG("Install_Delay_Days")},">0",{RG("Delivery_Delay_Days")},"<=0")&" of "&COUNTIF({RG("Install_Delay_Days")},">0")&" late installations followed an on-time delivery; top recorded reason: Customer site not ready ("&COUNTIFS({RG("Install_Delay_Days")},">0",{RG("Installation_Delay_Reason")},"Customer site not ready")&")"',
  "Observed (calculated) + recorded reasons", "Customer site readiness (civil/electrical/plumbing work) is not confirmed before the installation slot (hypothesis supported only by reason codes).",
  "Fixing installation scheduling would help even when the factory delivers on time."),
 ("Most overdue open orders are stuck at installation, and some are very old.",
  f'={KREF["overdue_orders"]}&" overdue orders: "&{KREF["overdue_install_stage"]}&" at installation, "&{KREF["overdue_delivery_stage"]}&" at delivery; average "&TEXT({KREF["avg_days_overdue"]},"0")&" days overdue; "&{KREF["pending_installations"]}&" delivered units await installation"',
  "Observed (calculated)", "No routine review of delivered-but-not-installed orders (hypothesis).",
  "Old overdue orders tie up value and risk cancellations (risk not measured)."),
 ("Production is the largest share of the order cycle.",
  f'="Avg production "&TEXT({KREF["avg_production"]},"0.0")&" d, transit "&TEXT({KREF["avg_transit"]},"0.0")&" d, delivery->install "&TEXT({KREF["avg_install_duration"]},"0.0")&" d; end-to-end cycle "&TEXT({KREF["avg_order_cycle"]},"0.0")&" d"',
  "Observed (calculated)", "Production time varies most for Custom and Commercial orders.",
  "Production capacity and material planning are the biggest levers on cycle time."),
 ("Orders to Other States are delivered on time less often.",
  f'="Other States "&TEXT(\'PIVOT ANALYSIS\'!$H${reg_rows["Other States"]},"0%")&" on time vs Bengaluru "&TEXT(\'PIVOT ANALYSIS\'!$H${reg_rows["Bengaluru"]},"0%")&" ("&\'PIVOT ANALYSIS\'!$G${reg_rows["Other States"]}&" timed deliveries - small sample)"',
  "Observed; small sample", "Longer and less predictable transport (hypothesis; built into the simulation).",
  "Transit allowances for distant regions may need review."),
 ("Delay reasons are missing on many late records.",
  f'={KREF["late_del_noreason"]}&" of "&{KREF["late_deliveries"]}&" late deliveries and "&{KREF["late_inst_noreason"]}&" of "&COUNTIF({RG("Install_Delay_Days")},">0")&" late installations have no reason"',
  "Observed (calculated)", "Reason entry is optional in the current process.",
  "Without reasons, root-cause work becomes guesswork."),
 ("The raw export needed cleaning before any KPI could be trusted.",
  '="3 duplicate rows removed; 48 changes or flags recorded in the cleaning log; "&' + KREF["dq_rows"] + '&" rows carry a data-quality flag"',
  "Observed (data-quality log)", "Free-text entry and manual exports.", "Untrusted numbers lead to wrong decisions."),
]
header_row(ws, 4, ["#", "Observed result", "Supporting metric (live formula)", "Evidence type", "Possible explanation (hypothesis unless stated)", "Why it matters"])
for i, (obs, met, ev, expl, why) in enumerate(FIND, start=1):
    rr = 4 + i
    for j, v in enumerate([f"F{i}", obs, met, ev, expl, why], start=1):
        c = ws.cell(rr, j, v); c.alignment = WRAP; c.font = font(9, j == 2); c.border = BOX
    ws.row_dimensions[rr].height = 54
RECS = [
 ("R1", "F5", "Daily 10-minute review of every overdue and delivered-not-installed order: owner, next step, revised date, customer informed (Y/N).", "Fewer orders forgotten after delivery; earlier customer communication.", "Overdue Orders; Average Days Overdue; Pending Installations", "High", "Low", "Immediate (week 1)", "Operations coordinator"),
 ("R2", "F3, F4", "Pre-installation readiness call/checklist 3-5 days before delivery (site, electrical, plumbing, access). Book the installation slot only after readiness is confirmed.", "Expected to reduce installations delayed by 'Customer site not ready'.", "On-Time Installation Rate; late installs with reason 'Customer site not ready'", "High", "Low", "Within 2 weeks", "Installation supervisor + customer service"),
 ("R3", "F2, F6", "Set promise dates by product category using actual production times (e.g. median + buffer) instead of one standard lead time.", "Promises that match capacity; higher on-time delivery for Commercial and Custom orders.", "On-time delivery % by category; gap between promised and actual lead time", "High", "Medium", "Within 1 month", "Sales operations + production planning"),
 ("R4", "F2", "Check component availability at order acceptance for Custom and Commercial orders; flag at-risk orders before a date is promised.", "Fewer long component-shortage delays.", "Late deliveries with reason 'Component shortage'; max delivery delay", "Medium", "Medium", "1-2 months", "Procurement / stores"),
 ("R5", "F8", "Make the delay reason a mandatory dropdown whenever a delivery or installation date passes its promise.", "Root causes become measurable.", "% of late records with a reason (target to discuss: 100%)", "Medium", "Low", "Within 2 weeks", "Operations analyst / data owner"),
 ("R6", "F9", "Improve order tracking: one tracker with standard status values, real date fields (no text dates), dispatch date and a duplicate-ID check.", "Clean data at source; less manual cleaning.", "Rows with data-quality flags per month (target: 0)", "Medium", "Low", "1 month", "Operations analyst"),
 ("R7", "F3, F5", "Installation scheduling: keep a shared technician calendar and plan capacity for busy months (only 4 'Technician availability' cases - a precaution, not a proven fix).", "Installations booked against real capacity.", "Install lead days (delivery->install); late installs with reason 'Technician availability'", "Medium", "Medium", "Before next peak season", "Installation supervisor"),
 ("R8", "F1-F8", "Weekly operations KPI report from this workbook: on-time delivery, on-time installation, overdue orders, backlog, reason completeness.", "Problems seen weekly instead of at month end.", "Report issued every week; actions closed on time", "Medium", "Low", "Start now", "Operations analyst"),
 ("R9", "F7", "Review transit allowance for Other States and agree pickup dates with transporters at dispatch.", "More reliable delivery promises outside Karnataka.", "On-time delivery % for Other States; average transit days", "Low", "Low", "1-2 months", "Dispatch / logistics"),
]
rr = 6 + len(FIND)
ws.cell(rr - 1, 1, "RECOMMENDATIONS - prioritised by impact, effort and urgency (all are proposals)").font = font(11, True, TEAL)
header_row(ws, rr, ["#", "Addresses", "Recommendation", "Expected benefit (not verified)", "Metric to monitor improvement", "Impact", "Effort", "Urgency", "Suggested owner (role)"])
for i, rec in enumerate(RECS, start=1):
    for j, v in enumerate(rec, start=1):
        c = ws.cell(rr + i, j, v); c.alignment = WRAP; c.font = font(9, j == 3); c.border = BOX
        if j == 6: c.font = font(9, True, RED if v == "High" else (AMBER if v == "Medium" else MID))
    ws.row_dimensions[rr + i].height = 52
for j, w in enumerate([6, 34, 46, 26, 40, 30, 9, 16, 22], start=1): ws.column_dimensions[CL(j)].width = w
ws.cell(rr + len(RECS) + 2, 1, "Limitations: simulated data; patterns such as longer custom-order production were assumptions built into the simulation, so they illustrate the method rather than reveal real causes. Category and region groups contain 17-49 timed deliveries each - treat small differences with caution.").font = font(9, i=True, c=RED)

# ======================= DATA QUALITY LOG =======================
ws = S["DATA QUALITY LOG"]
title(ws, "DATA QUALITY LOG - checks, rules and every change made", SYN)
ws["A4"] = "Part 1 - Live checks on CLEANED DATA (formulas)"; ws["A4"].font = font(11, True, TEAL)
header_row(ws, 5, ["#", "Check", "Result", "Expected", "Status"])
live = [
 ("Unique orders", f"=COUNTA({RG('Order_ID')})", "160", None),
 ("Duplicate Order IDs", f'=SUMPRODUCT(({RG("Order_ID")}<>"")*(COUNTIF({RG("Order_ID")},{RG("Order_ID")})>1))', "0", 0),
 ("Status outside Completed/Pending/Cancelled", f'=COUNTA({RG("Order_Status")})-COUNTIF({RG("Order_Status")},"Completed")-COUNTIF({RG("Order_Status")},"Pending")-COUNTIF({RG("Order_Status")},"Cancelled")', "0", 0),
 ("Category outside standard list", f'=COUNTA({RG("Product_Category")})-SUMPRODUCT(COUNTIF({RG("Product_Category")},LISTS!$A$3:$A$7))', "0", 0),
 ("Missing order date", f'=SUMPRODUCT(({RG("Order_ID")}<>"")*({RG("Order_Date")}=""))', "0", 0),
 ("Completed orders without delivery date", f'=COUNTIFS({RG("Order_Status")},"Completed",{RG("Actual_Delivery_Date")},"")', "1 (invalid date set blank)", None),
 ("Pending orders counted as installed", f'=COUNTIFS({RG("Order_Status")},"Pending",{RG("Installed")},1)', "0", 0),
 ("Cancelled orders inside delivery timing KPIs", f'=COUNTIFS({RG("Order_Status")},"Cancelled",{RG("Delivery_Delay_Days")},">-9999")', "0", 0),
 ("Negative production/transit/install durations used", f'=COUNTIF({RG("Production_Days")},"<0")+COUNTIF({RG("Transit_Days")},"<0")+COUNTIF({RG("Install_Lead_Days")},"<0")', "0", 0),
 ("Orders overdue although promise is after as-of date", f'=SUMPRODUCT(({RG("Overdue")}=1)*({RG("Delivered")}=0)*({RG("Expected_Delivery_Date")}>=AsOf))+SUMPRODUCT(({RG("Overdue")}=1)*({RG("Delivered")}=1)*({RG("Expected_Installation_Date")}>=AsOf))', "0", 0),
 ("Rows with a data-quality flag", f'=COUNTIFS({RG("DQ_Flag")},"<>OK",{RG("DQ_Flag")},"<>")', "9 (documented below)", None),
]
for i, (lab, f_, exp, ok) in enumerate(live, start=1):
    rr = 5 + i
    ws.cell(rr, 1, i); ws.cell(rr, 2, lab); ws.cell(rr, 3, f_); ws.cell(rr, 4, exp)
    ws.cell(rr, 5, f'=IF(C{rr}={ok},"OK","CHECK")' if ok is not None else "INFO")
    for j in range(1, 6): ws.cell(rr, j).border = BOX; ws.cell(rr, j).font = font(9, j == 5)
ws.conditional_formatting.add(f"E6:E{5+len(live)}", CellIsRule(operator="equal", formula=['"OK"'], font=Font(name=F, bold=True, color=GREEN)))
ws.conditional_formatting.add(f"E6:E{5+len(live)}", CellIsRule(operator="equal", formula=['"CHECK"'], font=Font(name=F, bold=True, color=RED)))
rr = 7 + len(live)
ws.cell(rr, 1, "Part 2 - Cleaning rules applied").font = font(11, True, TEAL); rr += 1
header_row(ws, rr, ["Rule", "Situation", "Action", "Why"]); rr += 1
RULES = [
 ("R1", "Spaces before/after Order ID", "Trimmed", "Prevents false duplicates and failed lookups"),
 ("R2", "Exact duplicate rows (double export)", "Removed second copy (3 rows)", "Each order counted once"),
 ("R3", "Spelling variants of status, category, customer type", "Mapped to one standard value; 'In Progress' = Pending", "COUNTIFS needs exact labels"),
 ("R4", "Dates as text, mixed formats, impossible date (31-02-2026)", "Converted to real dates; impossible date set blank and row flagged", "Never guess a date"),
 ("R5", "Order value blank or typed as 'Rs. 2,45,000'", "Text stripped to a number; blanks kept blank", "Value KPIs exclude blanks only"),
 ("R6", "Impossible sequence (installed before delivery; dispatched before order)", "Original dates kept; row excluded only from the affected duration/timing KPI", "Do not silently overwrite business data; confirm with owner"),
 ("R7", "Completed order with no installation date", "Status kept (system of record); excluded from installation timing; flagged", "Completion cannot be timed without a date"),
 ("R8", "Cancelled order that had been delivered (returned unit)", "Kept as Cancelled; excluded from delivery KPIs; not counted as a delivery failure", "Cancellation is not a delivery failure"),
 ("R9", "Late record with no delay reason", "Left blank, reported as a completeness gap", "Reasons are never invented"),
 ("R10", "Promise date after the as-of date", "Treated as 'not yet due', never overdue", "Future promises cannot be broken yet"),
]
for rule in RULES:
    for j, v in enumerate(rule, start=1):
        c = ws.cell(rr, j, v); c.font = font(9, j == 1); c.alignment = WRAP; c.border = BOX
    rr += 1
rr += 1
ws.cell(rr, 1, "Part 3 - Cleaning log (every changed or flagged cell)").font = font(11, True, TEAL); rr += 1
log = pd.read_csv(ROOT / "data/processed/cleaning_log.csv", dtype=str).fillna("")
header_row(ws, rr, ["Order ID", "Field", "Original value", "Cleaned value", "Rule applied"]); rr += 1
for rec in log.itertuples(index=False):
    for j, v in enumerate(rec, start=1):
        c = ws.cell(rr, j, v); c.font = font(8); c.border = BOX; c.alignment = WRAP
    rr += 1
for j, w in enumerate([12, 44, 26, 26, 70], start=1): ws.column_dimensions[CL(j)].width = w

# ======================= DATA DICTIONARY =======================
ws = S["DATA DICTIONARY"]
title(ws, "DATA DICTIONARY", SYN)
ws["A4"] = "RAW DATA columns (as exported)"; ws["A4"].font = font(11, True, TEAL)
header_row(ws, 5, ["Column", "Type", "Description", "Example"])
rr = 6
for rec in DD_RAW:
    for j, v in enumerate(rec, start=1):
        c = ws.cell(rr, j, v); c.font = font(9, j == 1); c.alignment = WRAP; c.border = BOX
    rr += 1
rr += 1
ws.cell(rr, 1, "CLEANED DATA columns").font = font(11, True, TEAL); rr += 1
header_row(ws, rr, ["Column", "Type", "Source", "Meaning", "Rule / formula", "Example"]); rr += 1
for rec in DD_CLEAN + DD_X:
    for j, v in enumerate(rec, start=1):
        c = ws.cell(rr, j, v); c.font = font(9, j == 1); c.alignment = WRAP; c.border = BOX
    rr += 1
for j, w in enumerate([28, 10, 14, 44, 56, 22], start=1): ws.column_dimensions[CL(j)].width = w

# ======================= START HERE =======================
ws = S["START HERE"]
ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 30; ws.column_dimensions["C"].width = 100
ws.merge_cells("B2:C2"); ws["B2"] = "Order Fulfilment & Installation Performance Analysis Dashboard"
ws["B2"].font = font(18, True, "FFFFFF"); ws["B2"].fill = fill(CHAR); ws.row_dimensions[2].height = 34
ws["B2"].alignment = Alignment(vertical="center", indent=1)
ws.merge_cells("B3:C3"); ws["B3"] = "Business Operations Analytics Case Study for Neo Heat Pump - Sai Heating and Cooling Solutions"
ws["B3"].font = font(11, True, "FFFFFF"); ws["B3"].fill = fill(TEAL); ws["B3"].alignment = Alignment(indent=1)
ws.merge_cells("B5:C5")
ws["B5"] = ("IMPORTANT - SIMULATED DATA: all 160 orders are synthetic records generated for a portfolio case study. "
            "They do not describe real orders, customers, prices or performance of the company. "
            "Findings illustrate an analysis method; recommendations are proposals only.")
ws["B5"].font = font(10, True, RED); ws["B5"].fill = fill("FDECEA"); ws["B5"].alignment = WRAP; ws.row_dimensions[5].height = 44
info = [
 ("What this is", "An operations-analysis workbook that turns raw order records into KPIs, a management dashboard, findings and recommendations for order fulfilment, delivery and installation."),
 ("Why it was built", "Management needs one reliable view of: how many orders are done, pending or cancelled; whether deliveries and installations keep their promised dates; which open orders are overdue; and where the process gets stuck."),
 ("Data period", "Orders booked 1-Jan-2026 to 30-Jun-2026. Report as-of date: 30-Jun-2026 (editable on KPI ANALYSIS, cell C4)."),
 ("Prepared by", "[Your Name] - portfolio project (not an employee report)."),
]
rr = 7
for a, b in info:
    ws.cell(rr, 2, a).font = font(10, True); ws.cell(rr, 3, b).font = font(10); ws.cell(rr, 3).alignment = WRAP
    ws.row_dimensions[rr].height = 30; rr += 1
rr += 1
ws.cell(rr, 2, "HOW TO NAVIGATE (click a sheet name)").font = font(11, True, TEAL); rr += 1
nav = [
 ("OPERATIONS DASHBOARD", "Start here if you have 2 minutes: KPI cards, 6 charts, overdue-order list. Yellow boxes filter everything."),
 ("KPI ANALYSIS", "Every KPI with formula, numerator, denominator, exclusions and interpretation."),
 ("PIVOT ANALYSIS", "Summary tables by product, month, customer type, region, stage and delay reason."),
 ("BUSINESS RECOMMENDATIONS", "Findings (with live supporting numbers) and 9 prioritised recommendations."),
 ("CLEANED DATA", "The validated dataset. Dark headers = inputs, teal headers = formula columns."),
 ("RAW DATA", "The original export, untouched, including its errors."),
 ("DATA QUALITY LOG", "Live quality checks, the cleaning rules and every cell that was changed or flagged."),
 ("DATA DICTIONARY", "Meaning of every column in RAW and CLEANED data."),
]
for nm, desc in nav:
    c = ws.cell(rr, 2, nm); c.hyperlink = f"#'{nm}'!A1"; c.font = Font(name=F, size=10, bold=True, color=TEAL, underline="single")
    ws.cell(rr, 3, desc).font = font(10); ws.cell(rr, 3).alignment = WRAP; rr += 1
rr += 1
ws.cell(rr, 2, "HOW TO READ THE DASHBOARD").font = font(11, True, TEAL); rr += 1
for a, b in [
 ("Order completion", "Completed (installed) orders ÷ non-cancelled orders. Recent orders that are not due yet still count as open, so June looks lower."),
 ("On-time delivery", "Deliveries on or before the promised date ÷ deliveries with valid dates. Undelivered and cancelled orders are not in this rate."),
 ("Installation completion", "Installations that were due by the as-of date and are done ÷ installations due."),
 ("Overdue", "An open order past the promise for the stage it is in. A promise after the as-of date is 'not yet due', never overdue."),
 ("Colours", "Teal = performance, red = late/overdue, amber = in progress, grey = cancelled/neutral."),
]:
    ws.cell(rr, 2, a).font = font(10, True); ws.cell(rr, 3, b).font = font(10); ws.cell(rr, 3).alignment = WRAP
    ws.row_dimensions[rr].height = 28; rr += 1
rr += 1
ws.cell(rr, 2, "TOOLS").font = font(11, True, TEAL); rr += 1
ws.cell(rr, 3, "Excel formulas (COUNTIFS, SUMIFS, AVERAGEIFS, INDEX/MATCH, LARGE, SUMPRODUCT), named cells, data validation dropdowns, conditional formatting, charts. "
              "The file was generated with Python (pandas, openpyxl) and recalculated in LibreOffice; open it in Excel and press Ctrl+Alt+F9 once to force a full recalculation. "
              "PivotTables/slicers are not included because the generator cannot create them - PIVOT ANALYSIS gives formula equivalents and 5 steps to add one.").alignment = WRAP
ws.cell(rr, 3).font = font(10); ws.row_dimensions[rr].height = 58

wb.active = wb.sheetnames.index("START HERE")
out = ROOT / "excel/operations_analysis.xlsx"
wb.save(out)
json.dump({"KREF": KREF, "PIV": PIV, "COL": COL}, open(ROOT / "scripts/_workbook_map.json", "w"), indent=1)
print("saved", out, "KPIs:", len(KREF))
